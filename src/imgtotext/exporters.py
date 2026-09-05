from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageOps

from imgtotext.models import ImageItem, ImageStatus


class ExportError(RuntimeError):
    pass


def completed_items(items: Iterable[ImageItem]) -> list[ImageItem]:
    return [item for item in items if item.status == ImageStatus.COMPLETED and item.description.strip()]


def export_json(items: Iterable[ImageItem], destination: Path) -> None:
    selected = completed_items(items)
    _require_results(selected)
    payload = {
        "application": "ImgToText",
        "exported_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "count": len(selected),
        "images": [item.as_export_dict(index) for index, item in enumerate(selected, start=1)],
    }
    destination.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def export_docx(items: Iterable[ImageItem], destination: Path) -> None:
    selected = completed_items(items)
    _require_results(selected)
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches
    except ImportError as exc:
        raise ExportError("Falta la dependencia python-docx") from exc
    document = Document()
    document.add_heading("Descripciones de imagenes", level=0)
    document.add_paragraph(f"Generado por ImgToText · {len(selected)} imagenes")
    for index, item in enumerate(selected, start=1):
        document.add_heading(f"{index}. {item.display_name}", level=1)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(item.path), width=Inches(5.8))
        document.add_paragraph(item.description)
        if item.visible_text:
            document.add_heading("Texto visible", level=2)
            document.add_paragraph(item.visible_text)
        if item.uncertainty:
            document.add_paragraph(f"Observacion: {item.uncertainty}")
        if index < len(selected):
            document.add_page_break()
    document.save(destination)


def export_pdf(items: Iterable[ImageItem], destination: Path) -> None:
    selected = completed_items(items)
    _require_results(selected)
    try:
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import Image as ReportImage
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:
        raise ExportError("Falta la dependencia reportlab") from exc

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ImageTitle", parent=styles["Heading1"], alignment=TA_CENTER, spaceAfter=12))
    story = [Paragraph("Descripciones de imagenes", styles["Title"]), Spacer(1, 0.4 * cm)]
    for index, item in enumerate(selected, start=1):
        story.append(Paragraph(_escape(f"{index}. {item.display_name}"), styles["ImageTitle"]))
        width, height = _display_size(item.path, max_width=16 * cm, max_height=14 * cm)
        story.append(ReportImage(str(item.path), width=width, height=height))
        story.append(Spacer(1, 0.5 * cm))
        story.append(Paragraph(_escape(item.description), styles["BodyText"]))
        if item.visible_text:
            story.append(Spacer(1, 0.25 * cm))
            story.append(Paragraph(f"<b>Texto visible:</b> {_escape(item.visible_text)}", styles["BodyText"]))
        if item.uncertainty:
            story.append(Spacer(1, 0.25 * cm))
            story.append(Paragraph(f"<b>Observacion:</b> {_escape(item.uncertainty)}", styles["BodyText"]))
        if index < len(selected):
            story.append(PageBreak())
    document = SimpleDocTemplate(
        str(destination),
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=1.8 * cm,
        bottomMargin=1.8 * cm,
        title="Descripciones de imagenes",
        author="ImgToText",
    )
    document.build(story)


def _display_size(path: Path, max_width: float, max_height: float) -> tuple[float, float]:
    with Image.open(path) as source:
        image = ImageOps.exif_transpose(source)
        width, height = image.size
    scale = min(max_width / width, max_height / height)
    return width * scale, height * scale


def _escape(value: str) -> str:
    return value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")


def _require_results(items: list[ImageItem]) -> None:
    if not items:
        raise ExportError("No hay descripciones completadas para exportar")
